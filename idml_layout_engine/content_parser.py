"""
Content Parser - Word Document Parsing
=======================================

This module handles parsing of Word documents (.docx) to extract
structured content blocks for layout assignment.

Word Document Structure (Expected):
- First paragraph = Section title
- Remaining paragraphs = Body text
- Quote paragraphs identified by specific rules

The parser produces ContentSection objects that contain:
- Title ContentBlock
- Body ContentBlocks (paragraphs and quotes)
- Associated image paths

Dependencies:
- python-docx: For reading .docx files

Usage:
    parser = ContentParser()
    sections = parser.parse_documents([
        "section1.docx",
        "section2.docx"
    ])
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import List, Dict, Optional, Tuple
from dataclasses import dataclass, field
import logging
import hashlib

# python-docx import - will be installed as dependency
try:
    from docx import Document as DocxDocument
    from docx.text.paragraph import Paragraph
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    DocxDocument = None
    Paragraph = None

from .models import (
    ContentBlock,
    ContentSection,
    ContentType,
)


logger = logging.getLogger(__name__)


# =============================================================================
# Quote Detection Configuration
# =============================================================================

# ============================================================================
# TODO: [INSERT FINAL QUOTE DETECTION RULE HERE]
# The following rules determine how quote paragraphs are identified.
# Update these based on your actual document conventions.
# ============================================================================

@dataclass
class QuoteDetectionConfig:
    """
    Configuration for detecting quote paragraphs.
    
    Quotes can be identified by:
    1. Style names (e.g., "Quote", "Pull Quote")
    2. Prefix patterns (e.g., starting with "> " or '"')
    3. Length heuristics (e.g., short paragraphs between 20-200 chars)
    4. Special markers (e.g., [QUOTE] prefix)
    
    Set the appropriate detection methods for your documents.
    """
    
    # Style-based detection
    quote_style_names: List[str] = field(default_factory=lambda: [
        # TODO: Add your quote style names
        # "Quote",
        # "PullQuote", 
        # "Blockquote",
    ])
    
    # Prefix-based detection
    quote_prefix_patterns: List[str] = field(default_factory=lambda: [
        # TODO: Add your quote prefix patterns
        # r'^>\s+',           # Markdown-style quote
        # r'^"\s*',           # Starts with quote mark
        # r'^\[QUOTE\]\s*',   # Tagged quote
    ])
    
    # Length-based detection (use with caution)
    use_length_heuristic: bool = False
    min_quote_length: int = 20
    max_quote_length: int = 200
    
    # Special markers
    quote_markers: List[str] = field(default_factory=lambda: [
        # "[QUOTE]",
        # "—",  # Em-dash often indicates attribution
    ])
    
    def matches_quote_style(self, style_name: Optional[str]) -> bool:
        """Check if a paragraph style indicates a quote."""
        if not style_name or not self.quote_style_names:
            return False
        style_lower = style_name.lower()
        return any(qs.lower() in style_lower for qs in self.quote_style_names)
    
    def matches_quote_prefix(self, text: str) -> bool:
        """Check if paragraph text starts with a quote prefix."""
        if not self.quote_prefix_patterns:
            return False
        for pattern in self.quote_prefix_patterns:
            if re.match(pattern, text):
                return True
        return False
    
    def matches_length_heuristic(self, text: str) -> bool:
        """Check if paragraph length suggests a quote."""
        if not self.use_length_heuristic:
            return False
        length = len(text.strip())
        return self.min_quote_length <= length <= self.max_quote_length
    
    def has_quote_marker(self, text: str) -> bool:
        """Check if text contains a quote marker."""
        if not self.quote_markers:
            return False
        return any(marker in text for marker in self.quote_markers)


# Default configuration instance
DEFAULT_QUOTE_CONFIG = QuoteDetectionConfig()


# =============================================================================
# Content Parser
# =============================================================================


class ContentParser:
    """
    Parser for extracting structured content from Word documents.
    
    Each .docx file is treated as one content section:
    - First paragraph = Title
    - Remaining paragraphs = Body (with quote detection)
    
    Usage:
        parser = ContentParser()
        sections = parser.parse_documents(["doc1.docx", "doc2.docx"])
        
        # Or with image mapping
        sections = parser.parse_documents(
            docx_paths,
            image_mapping={"Section Title": "path/to/image.jpg"}
        )
    """
    
    def __init__(
        self,
        quote_config: Optional[QuoteDetectionConfig] = None,
        strip_whitespace: bool = True,
        skip_empty_paragraphs: bool = True,
    ):
        """
        Initialize the content parser.
        
        Args:
            quote_config: Configuration for quote detection
            strip_whitespace: Whether to strip leading/trailing whitespace
            skip_empty_paragraphs: Whether to skip empty paragraphs
        """
        if not HAS_DOCX:
            raise ImportError(
                "python-docx is required for content parsing. "
                "Install it with: pip install python-docx"
            )
        
        self.quote_config = quote_config or DEFAULT_QUOTE_CONFIG
        self.strip_whitespace = strip_whitespace
        self.skip_empty_paragraphs = skip_empty_paragraphs
    
    def parse_documents(
        self,
        docx_paths: List[str | Path],
        image_mapping: Optional[Dict[str, str | List[str]]] = None,
    ) -> List[ContentSection]:
        """
        Parse multiple Word documents into content sections.
        
        Args:
            docx_paths: List of paths to .docx files
            image_mapping: Optional mapping of section titles to image paths
                          Can map to single path or list of paths
        
        Returns:
            List of ContentSection objects
        """
        image_mapping = image_mapping or {}
        sections = []
        
        for idx, docx_path in enumerate(docx_paths):
            path = Path(docx_path)
            if not path.exists():
                logger.warning(f"Document not found: {docx_path}")
                continue
            if not path.suffix.lower() == '.docx':
                logger.warning(f"Not a .docx file: {docx_path}")
                continue
            
            try:
                section = self._parse_single_document(path, idx, image_mapping)
                sections.append(section)
                logger.info(
                    f"Parsed {path.name}: title='{section.title.text[:50]}...', "
                    f"{len(section.body_blocks)} body blocks"
                )
            except Exception as e:
                logger.error(f"Error parsing {docx_path}: {e}")
                raise
        
        return sections
    
    def _parse_single_document(
        self,
        docx_path: Path,
        doc_index: int,
        image_mapping: Dict[str, str | List[str]],
    ) -> ContentSection:
        """
        Parse a single Word document into a ContentSection.
        
        Args:
            docx_path: Path to the .docx file
            doc_index: Index of this document (for ID generation)
            image_mapping: Title to image path mapping
            
        Returns:
            ContentSection object
        """
        doc = DocxDocument(str(docx_path))
        paragraphs = list(doc.paragraphs)
        
        if not paragraphs:
            raise ValueError(f"Document has no paragraphs: {docx_path}")
        
        # Generate section ID
        section_id = f"section_{doc_index}_{self._generate_id(docx_path)}"
        
        # First paragraph is the title
        title_para = paragraphs[0]
        title_text = self._extract_text(title_para)
        
        title_block = ContentBlock(
            content_id=f"{section_id}_title",
            content_type=ContentType.TITLE,
            text=title_text,
            section_title=title_text,
            sequence_index=0,
            source_file=str(docx_path),
        )
        
        # Remaining paragraphs are body content
        body_blocks = []
        sequence_idx = 1
        
        for para in paragraphs[1:]:
            text = self._extract_text(para)
            
            # Skip empty paragraphs if configured
            if self.skip_empty_paragraphs and not text.strip():
                continue
            
            # Determine if this is a quote
            content_type = self._classify_paragraph(para, text)
            
            block = ContentBlock(
                content_id=f"{section_id}_block_{sequence_idx}",
                content_type=content_type,
                text=text,
                section_title=title_text,
                sequence_index=sequence_idx,
                source_file=str(docx_path),
            )
            body_blocks.append(block)
            sequence_idx += 1
        
        # Get associated images
        # ====================================================================
        # TODO: [INSERT FINAL IMAGE–TITLE MAPPING RULE HERE]
        # The current implementation uses exact title matching.
        # Update this logic based on your actual mapping convention.
        # ====================================================================
        image_paths = self._get_images_for_title(title_text, image_mapping)
        
        return ContentSection(
            section_id=section_id,
            title=title_block,
            body_blocks=body_blocks,
            image_paths=image_paths,
            source_file=str(docx_path),
        )
    
    def _extract_text(self, paragraph: "Paragraph") -> str:
        """
        Extract text from a paragraph.
        
        Args:
            paragraph: python-docx Paragraph object
            
        Returns:
            Extracted text, optionally stripped
        """
        text = paragraph.text
        if self.strip_whitespace:
            text = text.strip()
        return text
    
    def _classify_paragraph(
        self, 
        paragraph: "Paragraph", 
        text: str
    ) -> ContentType:
        """
        Classify a paragraph as BODY or QUOTE.
        
        Uses the quote detection configuration to determine
        if a paragraph should be treated as a quote.
        
        Args:
            paragraph: python-docx Paragraph object
            text: Extracted text content
            
        Returns:
            ContentType.QUOTE or ContentType.BODY
        """
        config = self.quote_config
        
        # Check style name
        style_name = paragraph.style.name if paragraph.style else None
        if config.matches_quote_style(style_name):
            return ContentType.QUOTE
        
        # Check prefix patterns
        if config.matches_quote_prefix(text):
            return ContentType.QUOTE
        
        # Check for quote markers
        if config.has_quote_marker(text):
            return ContentType.QUOTE
        
        # Check length heuristic (use last, as it's least reliable)
        if config.matches_length_heuristic(text):
            return ContentType.QUOTE
        
        return ContentType.BODY
    
    def _get_images_for_title(
        self,
        title: str,
        image_mapping: Dict[str, str | List[str]],
    ) -> List[str]:
        """
        Get image paths for a section title.
        
        The mapping can use:
        1. Exact title match
        2. Normalized title match (lowercase, stripped)
        3. Partial match (title contains key)
        
        Args:
            title: The section title
            image_mapping: Title to image path(s) mapping
            
        Returns:
            List of image paths
        """
        # ====================================================================
        # TODO: [INSERT FINAL IMAGE–TITLE MAPPING RULE HERE]
        # Implement your specific mapping logic here.
        # Options:
        # - Exact match
        # - Case-insensitive match
        # - Fuzzy match
        # - Regex match
        # - Filename convention (e.g., "title_image.jpg")
        # ====================================================================
        
        images = []
        
        # Try exact match first
        if title in image_mapping:
            value = image_mapping[title]
            if isinstance(value, list):
                images.extend(value)
            else:
                images.append(value)
            return images
        
        # Try normalized match (lowercase, stripped)
        normalized_title = title.lower().strip()
        for key, value in image_mapping.items():
            if key.lower().strip() == normalized_title:
                if isinstance(value, list):
                    images.extend(value)
                else:
                    images.append(value)
                return images
        
        # Try partial match (title contains key)
        for key, value in image_mapping.items():
            if key.lower() in normalized_title or normalized_title in key.lower():
                if isinstance(value, list):
                    images.extend(value)
                else:
                    images.append(value)
                # Don't return - might have multiple partial matches
        
        return images
    
    def _generate_id(self, path: Path) -> str:
        """Generate a short unique ID from a file path."""
        hash_input = str(path.absolute()).encode()
        return hashlib.md5(hash_input).hexdigest()[:8]


# =============================================================================
# Utility Functions
# =============================================================================


def estimate_paragraph_units(text: str, chars_per_unit: int = 500) -> float:
    """
    Estimate how many paragraph units a text block occupies.
    
    This is a rough estimate based on character count.
    A "paragraph unit" is approximately one standard paragraph
    in the layout (about 3-5 lines of body text).
    
    Args:
        text: The text content
        chars_per_unit: Approximate characters per paragraph unit
        
    Returns:
        Estimated paragraph units (minimum 1.0)
    """
    if not text:
        return 1.0
    
    char_count = len(text.strip())
    units = max(1.0, char_count / chars_per_unit)
    
    return units


def validate_docx_structure(docx_path: str | Path) -> Dict[str, any]:
    """
    Validate that a Word document meets expected structure.
    
    Returns a dictionary with validation results:
    - is_valid: bool
    - errors: List of error messages
    - warnings: List of warning messages
    - paragraph_count: Number of paragraphs
    - has_title: Whether first paragraph looks like a title
    """
    result = {
        'is_valid': True,
        'errors': [],
        'warnings': [],
        'paragraph_count': 0,
        'has_title': False,
    }
    
    if not HAS_DOCX:
        result['is_valid'] = False
        result['errors'].append("python-docx not installed")
        return result
    
    path = Path(docx_path)
    if not path.exists():
        result['is_valid'] = False
        result['errors'].append(f"File not found: {docx_path}")
        return result
    
    if not path.suffix.lower() == '.docx':
        result['is_valid'] = False
        result['errors'].append(f"Not a .docx file: {docx_path}")
        return result
    
    try:
        doc = DocxDocument(str(path))
        paragraphs = list(doc.paragraphs)
        result['paragraph_count'] = len(paragraphs)
        
        if len(paragraphs) == 0:
            result['is_valid'] = False
            result['errors'].append("Document has no paragraphs")
            return result
        
        # Check first paragraph (title)
        title = paragraphs[0].text.strip()
        if title:
            result['has_title'] = True
            if len(title) > 200:
                result['warnings'].append(
                    f"Title is very long ({len(title)} chars), may not fit in title frame"
                )
        else:
            result['is_valid'] = False
            result['errors'].append("First paragraph (title) is empty")
        
        # Check for body content
        body_paragraphs = [p for p in paragraphs[1:] if p.text.strip()]
        if len(body_paragraphs) == 0:
            result['warnings'].append("Document has no body paragraphs")
        
        # Check for very long paragraphs
        for idx, para in enumerate(paragraphs):
            if len(para.text) > 3000:
                result['warnings'].append(
                    f"Paragraph {idx} is very long ({len(para.text)} chars)"
                )
        
    except Exception as e:
        result['is_valid'] = False
        result['errors'].append(f"Error reading document: {e}")
    
    return result
